#!/usr/bin/env python3
"""Describe the STRUCTURE of Excel files without printing any cell values.

Prints, per sheet: row and column counts, each column's header, how full it is,
its value type and typical text length. Nothing from inside the cells is shown,
so the output is safe to share when the data itself is confidential.
(Column headers ARE shown; check them before sharing.)

Setup:  pip install openpyxl
Run:    python3 inspect_excel.py /path/to/folder_or_file.xlsx
"""
import sys
from pathlib import Path

from openpyxl import load_workbook


def kind(v):
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, (int, float)):
        return "number"
    if isinstance(v, str):
        return "text"
    return type(v).__name__          # datetime, time, ...


def describe(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    print(f"\n=== {path.name} ===")
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        rows = [r for r in rows if any(c is not None and str(c).strip() for c in r)]
        if not rows:
            print(f"  [sheet] {ws.title}: empty")
            continue
        header, body = rows[0], rows[1:]
        print(f"  [sheet] {ws.title}: {len(body)} data rows, {len(header)} columns")
        for i, name in enumerate(header):
            vals = [r[i] for r in body if i < len(r) and r[i] is not None and str(r[i]).strip()]
            if not vals:
                print(f"    - {name!s:<32} empty")
                continue
            kinds = sorted({kind(v) for v in vals})
            texts = [len(v) for v in vals if isinstance(v, str)]
            extra = f", text length avg {sum(texts) // len(texts)} / max {max(texts)} chars" if texts else ""
            print(f"    - {name!s:<32} {100 * len(vals) // len(body):>3}% filled, "
                  f"{len(set(map(str, vals)))} distinct, {'/'.join(kinds)}{extra}")


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    target = Path(sys.argv[1])
    files = [target] if target.is_file() else sorted(
        p for p in target.rglob("*") if p.suffix.lower() in {".xlsx", ".xlsm"} and not p.name.startswith("~$"))
    if not files:
        sys.exit(f"No .xlsx files found under {target}")
    for f in files:
        try:
            describe(f)
        except Exception as e:                      # keep going on a broken file
            print(f"\n=== {f.name} === could not read: {e}")


if __name__ == "__main__":
    main()