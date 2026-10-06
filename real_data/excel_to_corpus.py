#!/usr/bin/env python3
"""Turn the HRD Excel files into a retrieval test set, entirely on this machine.

Reads the workbooks as they are (they are never modified) and writes:
  real/corpus.jsonl   one chunk per student row, plus one "card" chunk per sheet
                      describing what the sheet contains
  real/queries.jsonl  automatically generated questions with the expected chunk,
                      covering the seven query types

Layout it expects (found in every readable sheet of the sample files):
  title lines in column A, then a header block, then one row per student whose
  first cell is a running number and second cell is the name.

Nothing is sent anywhere and no cell values are printed; the summary shows
counts only. Add --show-headers to print the column names it detected.

Setup:  pip install openpyxl
Run:    python3 excel_to_corpus.py docs
Then:   python3 benchmark.py --corpus real/corpus.jsonl --queries real/queries.jsonl \
            --out results_real --models tfidf-baseline bge-small-en gte-modernbert arctic-l-v2 bge-m3
"""
import argparse
import datetime as dt
import json
import random
import re
from pathlib import Path

from openpyxl import load_workbook

# Invented names for no-answer questions; any that happen to exist in the data are skipped.
FAKE_NAMES = ["Zebulon Quartermain", "Ottoline Fairweather", "Barnaby Thistlewood", "Isolde Ravenscroft",
              "Leopold Ashgrove", "Wilhelmina Stroud", "Cornelius Pemberton", "Marguerite Holloway",
              "Thaddeus Blackwood", "Henrietta Castellane", "Ignatius Merriweather", "Rosalind Featherstone"]


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def is_text(v):
    return isinstance(v, str) and v.strip() != ""


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else f"{v:.2f}".rstrip("0").rstrip(".")
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    return re.sub(r"\s+", " ", str(v)).strip()


def col_letter(i):
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def read_sheet(ws):
    """Return dict(title, headers, rows[(excel_row, {header: text})], notes) or None."""
    raw = [list(r) for r in ws.iter_rows(values_only=True)]
    if not raw:
        return None
    ncols = max(len(r) for r in raw)
    raw = [r + [None] * (ncols - len(r)) for r in raw]

    # Header block starts at the first row with several text cells ...
    h = next((i for i, r in enumerate(raw[:30]) if sum(is_text(c) for c in r) >= 4), None)
    if h is None:
        return None
    # ... and ends where the student rows begin (running number + name).
    d = next((i for i in range(h + 1, len(raw)) if is_num(raw[i][0]) and is_text(raw[i][1])), None)
    if d is None:
        return None

    # Spread merged header cells over the columns/rows they cover.
    filled = [r[:] for r in raw]
    for rng in ws.merged_cells.ranges:
        top = raw[rng.min_row - 1][rng.min_col - 1]
        for rr in range(rng.min_row - 1, min(rng.max_row, len(raw))):
            for cc in range(rng.min_col - 1, min(rng.max_col, ncols)):
                filled[rr][cc] = top

    headers, seen = [], {}
    for c in range(ncols):
        parts = []
        for r in range(h, d):
            t = fmt(filled[r][c])
            if t and (not parts or parts[-1] != t):
                parts.append(t)
        name = " ".join(parts) or f"Column {col_letter(c)}"
        seen[name] = seen.get(name, 0) + 1
        headers.append(name if seen[name] == 1 else f"{name} #{seen[name]}")

    title = " | ".join(fmt(r[0]) for r in raw[:h] if is_text(r[0]))
    rows, notes = [], []
    for i in range(d, len(raw)):
        r = raw[i]
        cells = {headers[c]: fmt(r[c]) for c in range(ncols) if fmt(r[c])}
        if not cells:
            continue
        if is_num(r[0]) and is_text(r[1]):
            rows.append((i + 1, cells))
        else:
            notes.append("; ".join(cells.values()))
    used = [hd for c, hd in enumerate(headers) if any(hd in cells for _, cells in rows)]
    return dict(title=title, headers=used, name_col=headers[1], rows=rows, notes=notes)


def build(target, show_headers):
    files = [target] if target.is_file() else sorted(
        p for p in target.rglob("*") if p.suffix.lower() in {".xlsx", ".xlsm"} and not p.name.startswith("~$"))
    chunks, sheets = [], []
    for f in files:
        try:
            with open(f, "rb") as fh:
                if fh.read(2) != b"PK":
                    print(f"SKIP {f.name}: not an .xlsx workbook")
                    continue
            wb = load_workbook(f, data_only=True)
        except Exception as e:
            print(f"SKIP {f.name}: {type(e).__name__}: {e}")
            continue
        doc = f.stem.strip()
        for ws in wb.worksheets:
            sheet = ws.title.strip()
            info = read_sheet(ws)
            if not info or not info["rows"]:
                print(f"SKIP {f.name} / {sheet}: no student rows found")
                continue
            base = f"{doc}#{sheet}"
            card = (f"{info['title']}\nThis sheet has {len(info['rows'])} student rows. "
                    f"Columns: {', '.join(info['headers'])}.")
            if info["notes"]:
                card += "\nNotes: " + " / ".join(info["notes"])
            chunks.append(dict(chunk_id=f"{base}#card", doc_id=doc, doc_title=doc, section=sheet, text=card))
            for excel_row, cells in info["rows"]:
                text = "; ".join(f"{k}: {v}" for k, v in cells.items())
                chunks.append(dict(chunk_id=f"{base}#r{excel_row}", doc_id=doc, doc_title=doc,
                                   section=sheet, text=text))
            sheets.append(dict(doc=doc, sheet=sheet, base=base, **info))
            print(f"OK   {f.name} / {sheet}: {len(info['rows'])} student rows, "
                  f"{len(info['headers'])} columns in use, {len(info['notes'])} note lines")
            if show_headers:
                print("       columns: " + " | ".join(info["headers"]))
    return chunks, sheets


def make_queries(sheets, per_type, seed):
    rng = random.Random(seed)
    recs = []                                              # one record per student row
    for s in sheets:
        for excel_row, cells in s["rows"]:
            name = cells.get(s["name_col"], "")
            if not name:
                continue
            value_cols = [k for k, v in cells.items() if k != s["name_col"] and re.fullmatch(r"-?\d+(\.\d+)?", v)]
            recs.append(dict(name=name, key=re.sub(r"\s+", " ", name.lower()), doc=s["doc"], sheet=s["sheet"],
                             cid=f"{s['base']}#r{excel_row}", cols=value_cols[1:] or value_cols))
    by_name = {}
    for r in recs:
        by_name.setdefault(r["key"], []).append(r)
    multi = [r for r in recs if len(by_name[r["key"]]) > 1]
    real = set(by_name)

    def pick(pool, n):
        pool = list(pool)
        rng.shuffle(pool)
        return pool[:n]

    out = []

    def add(qtype, text, expected, note=""):
        out.append(dict(query_id=f"R{len(out) + 1:03d}", type=qtype, query=text,
                        expected_chunks=expected, note=note))

    for r in pick([r for r in recs if r["cols"]], per_type):
        add("exact", f"What is {r['name']}'s {rng.choice(r['cols'])} in the {r['sheet']} sheet of {r['doc']}?",
            [r["cid"]], "Names the student, column, sheet and file.")
    para = ["How did {name} do in {sheet}?", "Show me the record for {name} from {sheet}.",
            "I need the details recorded for {name} under {sheet}.", "Look up {name} in the {sheet} list."]
    for r in pick(recs, per_type):
        add("paraphrased", rng.choice(para).format(**r), [r["cid"]], "No column or file name given.")
    for r in pick(recs, per_type):
        add("short", f"{r['name']} {r['sheet']}", [r["cid"]], "Name and sheet only.")
    for r in pick([r for r in recs if len(r["cols"]) >= 2], per_type):
        c1, c2 = rng.sample(r["cols"], 2)
        add("long", f"I am going through the file {r['doc']} and need to check one trainee before a meeting. "
                    f"Can you find the row for {r['name']} in the {r['sheet']} sheet? I mainly need {c1} and "
                    f"{c2}, but please bring back the whole record.", [r["cid"]], "Long request with extra detail.")
    for s in pick(sheets, per_type):
        add("technical", f"Which columns are recorded in the {s['sheet']} sheet of {s['doc']}?",
            [f"{s['base']}#card"], "Asks about the sheet's structure; expects the sheet card.")
    for r in pick(multi, per_type):
        others = len(by_name[r["key"]]) - 1
        add("similar-document", f"{r['name']}'s result in {r['sheet']} ({r['doc']})", [r["cid"]],
            f"The same student also appears in {others} other sheet(s).")
    fakes = [n for n in FAKE_NAMES if n.lower() not in real]
    for i, n in enumerate(pick(fakes, per_type)):
        s = rng.choice(sheets)
        add("no-answer", f"What did {n} score in {s['sheet']}?" if i % 2 else f"Find the record for {n} in {s['doc']}.",
            [], "Invented name that is not in the data.")
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("target", help="folder (or single .xlsx file) to read")
    ap.add_argument("--out", default="real", help="output folder (default: real)")
    ap.add_argument("--per-type", type=int, default=10, help="questions per query type (default 10)")
    ap.add_argument("--seed", type=int, default=7)
    ap.add_argument("--show-headers", action="store_true", help="print the detected column names")
    a = ap.parse_args()

    chunks, sheets = build(Path(a.target), a.show_headers)
    if not chunks:
        raise SystemExit("No usable sheets found.")
    ids = [c["chunk_id"] for c in chunks]
    assert len(ids) == len(set(ids)), "duplicate chunk ids"
    queries = make_queries(sheets, a.per_type, a.seed)
    assert all(e in set(ids) for q in queries for e in q["expected_chunks"])

    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    with open(out / "corpus.jsonl", "w", encoding="utf-8") as f:
        for c in chunks:
            f.write(json.dumps(c, ensure_ascii=False) + "\n")
    with open(out / "queries.jsonl", "w", encoding="utf-8") as f:
        for q in queries:
            f.write(json.dumps(q, ensure_ascii=False) + "\n")

    counts = {}
    for q in queries:
        counts[q["type"]] = counts.get(q["type"], 0) + 1
    words = sorted(len(c["text"].split()) for c in chunks)
    print(f"\n{len(chunks)} chunks from {len(sheets)} sheets -> {out / 'corpus.jsonl'}")
    print(f"chunk length in words: median {words[len(words) // 2]}, longest {words[-1]}")
    print(f"{len(queries)} questions -> {out / 'queries.jsonl'}: " + ", ".join(f"{k} {v}" for k, v in counts.items()))


if __name__ == "__main__":
    main()